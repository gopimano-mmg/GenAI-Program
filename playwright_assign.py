from playwright.sync_api import sync_playwright
from openpyxl import load_workbook, Workbook
from datetime import datetime
from urllib.parse import quote
import json
import random
import time
import os

CONTACTS_FILE = "contacts.xlsx"
DATE_TEXT = datetime.now().strftime("%Y-%m-%d")
JSON_REPORT = f"whatsapp_report_{DATE_TEXT}.json"
EXCEL_REPORT = f"whatsapp_report_{DATE_TEXT}.xlsx"
SCREENSHOT_FOLDER = "screenshots"

os.makedirs(SCREENSHOT_FOLDER, exist_ok=True)


def random_delay():
    time.sleep(random.randint(2, 5))


def read_contacts():
    workbook = load_workbook(CONTACTS_FILE)
    sheet = workbook.active

    contacts = []

    for row in sheet.iter_rows(min_row=2, values_only=True):
        name, phone, message = row

        if not name or not phone:
            continue

        if not message:
            message = "Hello {name}, this is a test message from my Playwright bot."

        contacts.append({
            "name": str(name),
            "phone": str(phone),
            "message": str(message)
        })

    return contacts


def save_reports(results):
    with open(JSON_REPORT, "w", encoding="utf-8") as file:
        json.dump(results, file, indent=4, ensure_ascii=False)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "WhatsApp Report"

    sheet.append([
        "Name",
        "Phone",
        "Message",
        "Status",
        "Error",
        "Screenshot",
        "Last 3 Messages"
    ])

    for item in results:
        sheet.append([
            item.get("name"),
            item.get("phone"),
            item.get("message"),
            item.get("status"),
            item.get("error"),
            item.get("screenshot"),
            " | ".join(item.get("last_messages", []))
        ])

    workbook.save(EXCEL_REPORT)


def run_bot():
    contacts = read_contacts()
    results = []

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=False, slow_mo=300)
        context = browser.new_context()
        page = context.new_page()

        page.goto("https://web.whatsapp.com")

        print("Scan the QR code if WhatsApp is not logged in.")
        print("Waiting for WhatsApp Web to load...")

        page.wait_for_timeout(60000)

        for contact in contacts:
            name = contact["name"]
            phone = contact["phone"]
            message = contact["message"].replace("{name}", name)

            result = {
                "name": name,
                "phone": phone,
                "message": message,
                "status": "failed",
                "error": "",
                "screenshot": "",
                "last_messages": []
            }

            try:
                print(f"Sending message to {name} - {phone}")

                clean_phone = phone.replace("+", "").replace(" ", "").replace("-", "")
                encoded_message = quote(message)

                url = f"https://web.whatsapp.com/send?phone={clean_phone}&text={encoded_message}"
                page.goto(url)

                page.wait_for_timeout(15000)
                random_delay()

                message_box = page.locator("div[contenteditable='true']").last
                message_box.wait_for(timeout=60000)

                page.keyboard.press("Enter")

                page.wait_for_timeout(5000)
                random_delay()

                screenshot_path = os.path.join(
                    SCREENSHOT_FOLDER,
                    f"{name}_{DATE_TEXT}.png"
                )

                page.screenshot(path=screenshot_path, full_page=True)

                result["status"] = "sent"
                result["screenshot"] = screenshot_path

                messages = page.locator(
                    "div.message-in span.selectable-text, div.message-out span.selectable-text"
                )

                count = messages.count()
                last_messages = []
                start_index = max(0, count - 3)

                for i in range(start_index, count):
                    try:
                        text = messages.nth(i).inner_text()
                        last_messages.append(text)
                    except:
                        pass

                result["last_messages"] = last_messages

            except Exception as error:
                result["error"] = str(error)
                print(f"Error for {name}: {error}")

            results.append(result)
            random_delay()

        save_reports(results)

        print("Bot completed.")
        print("JSON report:", JSON_REPORT)
        print("Excel report:", EXCEL_REPORT)

        browser.close()


if __name__ == "__main__":
    run_bot()